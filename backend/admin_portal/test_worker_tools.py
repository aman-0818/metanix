"""Real executable checks; run in the built document-worker image."""
import os
import shutil
import tempfile
from unittest import skipUnless
from django.test import SimpleTestCase
from .services.pdf_operations import operate, office_to_pdf
from .services.document_converter import convert_document, cleanup_temp_file


class WorkerExecutableTests(SimpleTestCase):
    @skipUnless(shutil.which('soffice'), 'LibreOffice is installed in the worker image')
    def test_office_pdf_word_round_trip(self):
        from docx import Document
        import fitz
        with tempfile.TemporaryDirectory() as directory:
            source = os.path.join(directory, 'policy.docx')
            doc = Document()
            doc.add_heading('Travel policy', 1)
            doc.add_paragraph('Employees may claim 123 USD for approved travel.')
            doc.save(source)
            pdf = office_to_pdf(source, directory)
            with fitz.open(pdf) as result:
                self.assertIn('123 USD', ''.join(page.get_text() for page in result))
            success, converted, message = convert_document(pdf, 'docx')
            self.assertTrue(success, message)
            try:
                result = Document(converted)
                text = '\n'.join(p.text for p in result.paragraphs)
                text += '\n'.join(cell.text for table in result.tables for row in table.rows for cell in row.cells)
                self.assertIn('123 USD', text)
            finally:
                cleanup_temp_file(converted)

    @skipUnless(shutil.which('ocrmypdf'), 'OCRmyPDF is installed in the worker image')
    def test_scanned_pdf_becomes_searchable(self):
        import fitz
        with tempfile.TemporaryDirectory() as directory:
            source = os.path.join(directory, 'scan.pdf')
            with fitz.open() as original:
                page = original.new_page()
                page.insert_text((60, 120), 'Travel policy', fontsize=36)
                page.insert_text((60, 200), 'Approved reimbursement 123 USD', fontsize=24)
                image = page.get_pixmap(dpi=200).tobytes('png')
            with fitz.open() as scanned:
                page = scanned.new_page()
                page.insert_image(page.rect, stream=image)
                self.assertEqual(page.get_text().strip(), '')
                scanned.save(source)
            output = operate([source], 'ocr', {}, lambda *args: None)
            try:
                with fitz.open(output) as result:
                    text = ' '.join(page.get_text() for page in result)
                    self.assertIn('Travel', text)
                    self.assertIn('123', text)
            finally:
                shutil.rmtree(os.path.dirname(output))
